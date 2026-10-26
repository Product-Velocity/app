"""Turns docs/reference/Roadmap_v4_field_dictionary.xlsx into packages/domain/src/field-dictionary.json.

Run from the repo root:  python3 tools/field-dictionary/export.py
Needs: pip install openpyxl. The JSON is committed; re-run when the workbook changes.
"""
import csv
import json
import re
import openpyxl

SRC = "docs/reference/Roadmap_v4_field_dictionary.xlsx"
OUT = "packages/domain/src/field-dictionary.json"
VIEWS = ["Prioritisation", "Planning", "Delivery", "Go-to-market", "Outcomes", "Leadership", "Discovery", "History"]

CSV_SRC = "docs/reference/item-fields.csv"
csv_by_key = {r["key"]: r for r in csv.DictReader(open(CSV_SRC, encoding="utf-8"))}

wb = openpyxl.load_workbook(SRC, data_only=True)
dict_rows = list(wb["Field dictionary v4"].iter_rows(values_only=True))
dh = list(dict_rows[0])
by_key = {}
for r in dict_rows[1:]:
    if r[0] is None:
        continue
    d = dict(zip(dh, r))
    by_key[d["Field key (proposed)"]] = d

grid_rows = list(wb["Grid columns and views"].iter_rows(values_only=True))
gh = list(grid_rows[0])
out = []
for r in grid_rows[1:]:
    g = dict(zip(gh, r))
    key = g["Field key"]
    if key is None:  # blank and note rows; rows added later have no number on this sheet
        continue
    d = by_key.get(key, {})
    allowed = d.get("Allowed values")
    values = [v.strip() for v in re.split(r"\s·\s", str(allowed))] if allowed and str(allowed) != "None" and "·" in str(allowed) else []
    out.append({
        "n": str(d.get("#") or g["#"]),
        "storage": csv_by_key[key]["storage"],
        "phase": int(csv_by_key[key]["ships_in_phase"]),
        "name": g["Field"],
        "key": key,
        "group": g["Group"],
        "type": g["Type"],
        "editing": g["Editing in grid"],
        "customerSafe": g["Customer-safe reports"] == "Y",
        "views": [v for v in VIEWS if g.get(v) == "Y"],
        "allowed": values,
        "allowedText": None if allowed in (None, "None") else str(allowed),
        "origin": d.get("Origin"),
        "githubSync": d.get("GitHub sync"),
        "filledWhen": d.get("Filled when"),
        "description": d.get("Description"),
    })

assert len(out) == len(csv_by_key), len(out)
assert {f["key"] for f in out} == set(csv_by_key), "workbook and CSV list different keys"
assert len({f["key"] for f in out}) == len(out), "duplicate field keys"
with open(OUT, "w") as f:
    json.dump(out, f, indent=1, ensure_ascii=False)
    f.write("\n")
print(f"wrote {len(out)} fields to {OUT}")

# The sheet "Grid columns and views" as a CSV the tests compare the default views against.
GRID_CSV = "docs/reference/grid-views.csv"
with open(GRID_CSV, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["key", "group", "editing", "customer_safe"] + VIEWS)
    for r in grid_rows[1:]:
        if r[2] is None:  # blank and note rows (rows added later have no number on this sheet)
            continue
        g = dict(zip(gh, r))
        w.writerow([g["Field key"], g["Group"], g["Editing in grid"], "Y" if g["Customer-safe reports"] == "Y" else ""] + ["Y" if g.get(v) == "Y" else "" for v in VIEWS])
print(f"wrote {GRID_CSV}")
